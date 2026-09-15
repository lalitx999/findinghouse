import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse

def export_deals_to_excel(queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "รายงานดีลการขายและเช่า"

    # Header styling
    header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "รหัสดีล/ชื่อดีล", "ประเภทดีล", "ขั้นตอน Pipeline", "ชื่อลูกค้า (Lead)",
        "รหัสทรัพย์สิน", "ราคาตกลง (บาท)", "อัตราคอมมิชชัน (%)",
        "ยอดคอมมิชชันรวม (บาท)", "ส่วนแบ่งบริษัท (บาท)", "ส่วนแบ่งนายหน้า (บาท)",
        "นายหน้าผู้ดูแล", "วันที่ปิดดีล"
    ]

    ws.append(headers)

    # Style Header Row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Data Rows
    for row_idx, deal in enumerate(queryset, start=2):
        row = [
            deal.title,
            deal.get_deal_type_display(),
            deal.get_stage_display(),
            deal.lead.name if deal.lead else "-",
            deal.property.property_code if deal.property else "-",
            float(deal.agreed_price) if deal.agreed_price else 0,
            float(deal.commission_rate) if deal.commission_rate else 0,
            float(deal.commission_amount) if deal.commission_amount else 0,
            float(deal.company_share) if deal.company_share else 0,
            float(deal.agent_share) if deal.agent_share else 0,
            deal.agent.username if deal.agent else "-",
            deal.actual_close_date.strftime("%d/%m/%Y") if deal.actual_close_date else "-"
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            if col_num in [6, 8, 9, 10]:
                cell.number_format = '#,##0.00'
                cell.alignment = right_align
            elif col_num in [7]:
                cell.number_format = '0.00'
                cell.alignment = right_align
            else:
                cell.alignment = left_align

    # Auto adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="findinghouse_deals_report.xlsx"'
    wb.save(response)
    return response
