import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from django.http import HttpResponse

def export_properties_excel(queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "รายการอสังหาริมทรัพย์"

    # Styling setup
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Dark Slate
    header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Tahoma", size=10)
    bold_font = Font(name="Tahoma", size=10, bold=True)
    border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        'รหัสทรัพย์', 'ชื่อประกาศ', 'ประเภท', 'หมวดหมู่', 'ทำเลที่ตั้ง',
        'ราคา (บาท)', 'จำนวนห้องนอน', 'จำนวนห้องน้ำ', 'พื้นที่ใช้สอย/ดิน',
        'ผู้ขายตรง/เอเจนท์', 'เบอร์โทรศัพท์', 'สถานะ', 'วันที่ลงประกาศ'
    ]
    ws.append(headers)

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_price = 0
    for prop in queryset:
        price_val = float(prop.price) if prop.price else 0.0
        total_price += price_val
        seller_info = prop.seller_name if prop.is_direct_seller else (f"{prop.agent.first_name} {prop.agent.last_name}" if prop.agent else "Finding House")
        seller_phone = prop.seller_phone if prop.is_direct_seller else "02-123-4567"
        size_info = f"{prop.size_usable or prop.size_land or '-'} ตร.ม./ตร.ว."

        row = [
            prop.property_code,
            prop.title,
            "ขาย" if prop.property_type == 'SALE' else "เช่า",
            prop.category.name if prop.category else "-",
            prop.location.name if prop.location else "-",
            price_val,
            prop.bedrooms or 0,
            prop.bathrooms or 0,
            size_info,
            seller_info,
            seller_phone,
            prop.get_status_display(),
            prop.created_at.strftime('%Y-%m-%d %H:%M')
        ]
        ws.append(row)

    # Format data rows
    for row in range(2, ws.max_row + 1):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = cell_font
            cell.border = border
            if col == 6: # Price column
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col in [1, 3, 7, 8, 12, 13]:
                cell.alignment = Alignment(horizontal="center")

    # Summary Row
    summary_row = ws.max_row + 1
    ws.cell(row=summary_row, column=2, value="รวมมูลค่าทรัพย์สินทั้งหมด:").font = bold_font
    price_cell = ws.cell(row=summary_row, column=6, value=total_price)
    price_cell.font = bold_font
    price_cell.number_format = '#,##0.00'
    price_cell.alignment = Alignment(horizontal="right")

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="FindingHouse_Properties_Report.xlsx"'
    return response
